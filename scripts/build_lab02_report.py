"""Build the Ukrainian Word report for laboratory work 2."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "report-lr02.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def border_cells(table):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "4")
                edge.set(qn("w:color"), "D9D9D9")
                borders.append(edge)
            tc_pr.append(borders)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        cell.text = header
        shade(cell, "E9EFF5")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for p in table.rows[0].cells[0].paragraphs:
        p.paragraph_format.keep_with_next = True
    for item in rows:
        cells = table.add_row().cells
        for i, value in enumerate(item):
            cells[i].width = Cm(widths[i])
            cells[i].text = value
    border_cells(table)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def paragraph(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def labelled(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    return p


def shot(doc, number, description):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_margins(cell, 100, 130, 100, 130)
    shade(cell, "F5F7FA")
    border_cells(table)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"МІСЦЕ ДЛЯ СКРИНШОТА {number}\n")
    r.bold = True
    r.font.color.rgb = RGBColor(42, 73, 102)
    p.add_run(description)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(1.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15
for name, size, before, after in (("Title", 17, 0, 12), ("Heading 1", 12, 10, 5), ("Heading 2", 10.5, 7, 3)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_ppr = styles["Title"]._element.get_or_add_pPr()
title_borders = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "nil")
title_borders.append(bottom)
title_ppr.append(title_borders)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("ЛР 02  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.add_paragraph("Звіт до лабораторної роботи 2", style="Title")
paragraph(doc, "Валідація API та усунення SQL ін’єкції у навчальному трекері інцидентів")
labelled(doc, "ПІБ і група:", "[заповнити студенту]")
labelled(doc, "Дата перевірки:", "05.10.2026. Заявлений рівень: добрий («Добре»).")

doc.add_heading("1 Ідентифікація", level=1)
add_table(doc, ["Параметр", "Фактичний стан"], [
    ("Baseline", "Спільний варіант 2-A «Трекер інцидентів», тег ЛР 1 v0.1.0, base commit 1f62a06983e567d39652277da202dc61d89719db."),
    ("Scaffold", "lab-02-start-v1; квитанція .scaffolds/lab-02.json; base_commit і entry_parent збігаються з 1f62a06983e567d39652277da202dc61d89719db."),
    ("Git", "На час підготовки звіту HEAD 1aa07f3 у гілці lab/2-system. Потрібну гілку lab/2-input-sqli, vulnerable/fixed commits і тег v0.2.0 ще не створено."),
    ("DEL-01", "Приватний remote або bundle ще не передано. Не позначати цей пункт виконаним до перевірки історії й доступу викладача."),
], [3.3, 13.7])
paragraph(doc, "Початковий код ЛР 2 та результати перевірок зберігаються у проєкті SecureLab. Документ не приписує нездійснених Git-дій. Після їх виконання студент має доповнити цю таблицю справжніми hashes і даними передавання.")
shot(doc, "1", "Після оформлення Git: історія з двома окремими commits, назва гілки та тег v0.2.0. Не вставляти скриншот, який показує інший стан.")

doc.add_heading("2 Контракт створення та CP 01", level=1)
paragraph(doc, "POST /api/incidents приймає окремий CreateIncidentRequest. Сервер перевіряє значення до створення сутності; успішна відповідь 201 Created містить лише id, title, severity, status і createdAtUtc та заголовок Location.")
add_table(doc, ["Поле", "Перевірка"], [
    ("title", "Обов'язкове; не лише пробіли; максимум 160 символів до Trim; збереження після Trim."),
    ("description", "Обов'язкове; максимум 4000 символів до Trim. Для High і Critical після Trim потрібно щонайменше 40 символів."),
    ("severity", "Enum.TryParse разом з Enum.IsDefined; рядок «7» і невідоме значення повертають 400."),
    ("occurredAtUtc", "Обов'язковий DateTimeOffset; не далі ніж 5 хвилин після UTC now, зафіксованого на початку обробки; перед записом переводиться в UTC."),
    ("Керовані сервером", "id, ownerUserId, status, createdAtUtc та updatedAtUtc відсутні у request DTO і призначаються сервером."),
], [3.3, 13.7])
paragraph(doc, "Повторний title після Trim повертає 409, якщо наявний інцидент має статус New, Triaged, InProgress або Resolved. Порівняння чутливе до регістру; Closed повторне створення не блокує. Помилки полів повертаються як 400 application/problem+json з errors, предметний конфлікт — 409 application/problem+json.")
paragraph(doc, "CP-01. Клієнт може надіслати JSON поза браузером, тому обмеження елемента select не заміняє перевірку сервера. Enum.TryParse розбирає числовий рядок «7», але Enum.IsDefined відхиляє значення, якого немає у переліку. Тест із зайвими id, ownerUserId, status і createdAtUtc підтверджує серверні значення збереженого запису.")
shot(doc, "2", "За бажанням: 400 Problem Details для severity «7» з errors.severity і 201/GET, де статус New та власник Аліса Коваль. Перед знімком прибрати секрети й зайві журнали.")

doc.add_heading("3 Сценарій безпеки та CP 02 і CP 03", level=1)
paragraph(doc, "До виправлення MapLab02Endpoints складав SQL конкатенацією q і допускав довільний sortBy через гілку «_ => sortBy». Це дозволяло вводу змінювати структуру запиту. Перевірку виконано тільки на локальному навчальному стенді, методом GET без зміни даних.")
add_table(doc, ["Запит", "До виправлення", "Після виправлення"], [
    ("A  q=USB", "200, один запис з ID …005", "200, той самий запис"),
    ("B  q=zz-no-match", "200, []", "200, []"),
    ("C  контрольний ввід із методички", "200, усі п'ять seed-записів", "200, []"),
    ("D  q=O'Brien", "500 Problem Details", "200, лише запис …004"),
], [5.6, 5.7, 5.7])
paragraph(doc, "Фактичні відповіді: docs/evidence/search-before.txt і docs/evidence/search-after.txt. Скрипт scripts/check-lab02-search.ps1 лише повторює чотири HTTP-запити; самі сценарії також наведено в tests/http/lab-02-checks.http.")
paragraph(doc, "CP-02. У стані «до» СУБД отримувала SQL-текст, у який q вже було вставлено як частину виразу. Порожня вибірка B перетворилася на всі записи C; легітимний апостроф D порушив синтаксис. Заборона апострофа зламала б правильний пошук і не виправила б причину.")
shot(doc, "3", "Локальний стан «до»: поруч показати B як 200 [] та C як 200 із п'ятьма навчальними записами; за потреби D з 500. Приховати connection string, credentials і повний SQL.")
paragraph(doc, "Після виправлення пошук побудовано через LINQ і EF.Functions.ILike за Title та Description. Pattern передається параметром, а %, _ і зворотна коса риска екрануються для буквального пошуку. sortBy приймає лише createdAtUtc, severity і status; невідоме значення повертає 400 з errors.sortBy. Сортування має бізнес-ранги, другий ключ Id і обмеження Take(50).")
paragraph(doc, "CP-03. Той самий контрольний ввід C після виправлення повернув 200 []. Звичайний пошук і апостроф працюють. Автотест перевіряє точну множину результатів, а тест із 55 записами — порядок та межу 50. У журналі виправленого API умова ILIKE містить параметри pattern, а не введений текст у структурі SQL.")
shot(doc, "4", "Локальний стан «після»: C як 200 [], D як 200 з одним записом; за можливості поруч фрагмент журналу з параметрами pattern без секретів.")

checks_heading = doc.add_heading("4 Фактичні перевірки DEL 02", level=1)
checks_heading.paragraph_format.page_break_before = True
paragraph(doc, "Виконано dotnet test tests/SecureLab.Api.Tests/SecureLab.Api.Tests.csproj --configuration Release --no-restore. Результат: 49 пройдено, 0 помилок, 0 пропущено. Release-збірка API і тестів пройшла без попереджень та помилок. Перелік тестів: docs/evidence/tests-summary.txt.")
add_table(doc, ["ID", "Сценарій", "Фактично та доказ"], [
    ("T-01", "Коректний POST", "201, вузький response DTO, Location, Trim/UTC, серверні значення. Тест Create_UsesServerValues_NormalizesUtc_AndReturnsOnlyContract."),
    ("T-02", "Некоректний DTO", "400 з errors; null/порожні поля, «7», майбутня дата, межі довжин. Тести InvalidDto_ReturnsFieldProblem_WithoutInsert та LengthLimit_IsMeasuredBeforeTrim."),
    ("T-03", "Дублікат title", "409 для чотирьох активних статусів; Closed дозволяє 201; регістр враховано. Тест DuplicateTitle_ConflictsUnlessClosed_AndIsCaseSensitive."),
    ("S-01", "SQLi до виправлення", "B: []; C: усі 5 seed-записів. docs/evidence/search-before.txt."),
    ("S-02", "Повтор після fix", "C: 200 []. docs/evidence/search-after.txt; Search_Regression_ReturnsExactSet."),
    ("T-04", "Нормальний пошук", "USB: лише …005; O'Brien: лише …004. docs/evidence/search-after.txt."),
    ("T-05", "Невідомий sortBy", "price, created_at_utc, Severity → 400 errors.sortBy. Тест UnknownSort_ReturnsValidationProblem."),
    ("T-06", "Відсутній ID", "404 application/problem+json, заголовок «Інцидент не знайдено». Тест MissingIncident_UsesSafeProblemDetails."),
    ("T-09", "Опис 39/40", "High і Critical: 39 → 400 description, 40 → 201. Тест CrossField_ChecksTrimmedDescription."),
    ("A-01", "Огляд доступу до БД", "Небезпечний raw SQL пошуку усунуто; статичний службовий SQL у reset не залежить від вводу. Див. розділ 5."),
], [1.3, 3.8, 11.9])
shot(doc, "5", "Вікно тестів або термінал із підсумком 49 passed, 0 failed; дата й назва збірки мають відповідати фактичному запуску.")

access_heading = doc.add_heading("5 Огляд доступу до даних A 01", level=1)
access_heading.paragraph_format.page_break_before = True
paragraph(doc, "Пошук у вихідному коді охопив FromSql, ExecuteSql, SqlQuery, DbCommand і операції EF. Зовнішні значення обробляють endpoints та IncidentQueries; прямого складання SQL із ними після виправлення немає.")
add_table(doc, ["Точка", "Категорія та висновок"], [
    ("Lab02Endpoints GET search", "LINQ/ILike, параметризований pattern, allowlist сортування, Take(50)."),
    ("Lab02Endpoints POST", "AnyAsync і SaveChangesAsync; title як значення параметра; валідація перед записом."),
    ("IncidentQueries", "GetListAsync, GetDetailsAsync, GetSeveritySummaryAsync — LINQ із параметрами, DTO-проєкції."),
    ("DbSeeder і Lab02Seed", "Фіксовані навчальні ID та значення, AnyAsync і SaveChangesAsync."),
    ("DatabaseBootstrap", "MigrateAsync; статичний ExecuteSqlRawAsync TRUNCATE лише для явно дозволеного Development reset."),
    ("Program GET health", "CanConnectAsync без SQL із вводу користувача."),
    ("Тести", "ExecuteUpdateAsync/ExecuteDeleteAsync стосуються локальних fixtures і очищення власних тестових записів."),
], [5.4, 11.6])

doc.add_heading("6 Обмеження та висновок", level=1)
paragraph(doc, "Для заявленого рівня «Добре» реалізовано серверну валідацію, конфлікт title, захист пошуку й автоматичні перевірки. Вразливий пошук до виправлення розширював вибірку, після виправлення той самий ввід не змінює умову запиту.")
paragraph(doc, "AnyAsync перед вставленням не гарантує унікальність за одночасних POST; для production потрібне обмеження БД або транзакційне рішення. Пошкоджений JSON окремо не перевірявся; результати валідації стосуються синтаксично правильного JSON. Автентифікація, авторизація, rate limit та повна пагінація не входять до ЛР 2. Міграцію БД не додано.")
paragraph(doc, "Git-історію та подання DEL-01/DEL-02 ще треба завершити. Без vulnerable/fixed commits і тегу робота не відповідає повному комплекту здачі, навіть якщо код і тести пройшли.")

doc.add_heading("7 Декларація використання ШІ", level=1)
add_table(doc, ["Сервіс", "Задачі", "Авторська перевірка"], [
    ("Codex", "Аналіз методички й коду, реалізація валідації та безпечного пошуку, тести, підготовка звіту.", "Реальні HTTP-відповіді «до/після» та запуск 49 тестів зафіксовано локально. Студенту слід самостійно перевірити й пояснити код перед захистом."),
], [2.5, 6.9, 7.6])

doc.add_heading("8 Передавання", level=1)
paragraph(doc, "Після створення Git-історії заповнити фактичні hashes vulnerable/fixed commits, назву переданого bundle або URL погодженого приватного remote, перевірити доступ викладача. Завершений звіт зберегти окремим commit, тег v0.2.0 поставити на нього. Перевірити staged diff та історію на відсутність секретів. Лише після цього позначити подання в Moodle виконаним.")
shot(doc, "6", "Після завершення: git log із vulnerable/fixed commits, git show v0.2.0, перевірка bundle або доступу до remote та підтвердження подання в Moodle. Вставити тільки власні фактичні знімки.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
